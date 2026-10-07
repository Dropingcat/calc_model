# -*- coding: utf-8 -*-
"""Верификация ФР-13: H28, Табл4."""
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
p = r'C:\Users\Arhys\Desktop\каллибровка\фракционный состав\фракционка ГОСТ_2177.xlsx'
wb = openpyxl.load_workbook(p, data_only=False)
print('sheets:', wb.sheetnames)
ws = wb['Лист1']
print('H28 =', repr(ws['H28'].value)[:300])
print('H29 =', repr(ws['H29'].value))
print('K28 =', repr(ws['K28'].value))
print('K29 =', repr(ws['K29'].value))
if 'Табл4' in wb.sheetnames:
    t = wb['Табл4']
    print('\n--- Табл4: первые/последние строки ---')
    for r in range(1, 26):
        a = t.cell(row=r, column=1).value
        b = t.cell(row=r, column=2).value
        c = t.cell(row=r, column=3).value
        d = t.cell(row=r, column=4).value
        if a is not None or b is not None:
            print('r%d: кПа=%r мм=%r A=%r B=%r' % (r, a, b, c, d))
wb.close()