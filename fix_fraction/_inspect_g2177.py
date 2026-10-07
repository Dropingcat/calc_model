# -*- coding: utf-8 -*-
"""Инспекция ГОСТ_2177 для ФР-13."""
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')
p = r'C:\Users\Arhys\Desktop\каллибровка\фракционный состав\фракционка ГОСТ_2177.xlsx'
wb = openpyxl.load_workbook(p, data_only=False)

ws = wb['Лист1']
print('=== Лист1: E1, F26, G26, H28, H29, K28, K29 ===')
for cell in ['E1', 'F26', 'G26', 'H28', 'H29', 'K28', 'K29', 'C28', 'C29', 'C30']:
    print('%s = %r' % (cell, ws[cell].value))

print('\n=== Лист1: строка 28-30 контекст (A..M) ===')
for r in range(25, 32):
    row = []
    for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M']:
        v = ws['%s%d' % (col, r)].value
        if v is not None:
            row.append('%s%d=%r' % (col, r, v))
    print(' | '.join(row) if row else '  (пусто)')

print('\n=== Лист1: заголовки строк 1..31 (A) ===')
for r in range(1, 32):
    a = ws['A%d' % r].value
    if a is not None:
        print('A%d: %r' % (r, a))

wb.close()

wb2 = openpyxl.load_workbook(p, data_only=False)
wv = wb2['Ввод данных']
print('\n=== Ввод данных: C28, E28..E30, E1 (если есть) ===')
for cell in ['E1', 'C28', 'E28', 'E29', 'E30', 'D28', 'D29', 'D30']:
    print('%s = %r' % (cell, wv[cell].value))
# строка 28-30 контекст
for r in range(26, 32):
    row = []
    for col in ['A', 'B', 'C', 'D', 'E']:
        v = wv['%s%d' % (col, r)].value
        if v is not None:
            row.append('%s%d=%r' % (col, r, v))
    print('Ввод строка %d: %s' % (r, ' | '.join(row) if row else '(пусто)'))
wb2.close()

wb3 = openpyxl.load_workbook(p, data_only=False)
wo = wb3['Вывод']
print('\n=== Вывод: E28..E30 и окрестности ===')
for r in range(25, 32):
    row = []
    for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
        v = wo['%s%d' % (col, r)].value
        if v is not None:
            row.append('%s%d=%r' % (col, r, v))
    print('Вывод строка %d: %s' % (r, ' | '.join(row) if row else '(пусто)'))
wb3.close()