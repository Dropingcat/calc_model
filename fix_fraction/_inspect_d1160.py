# -*- coding: utf-8 -*-
"""Инспекция MATCH-индексов в D1160 Лист2."""
import openpyxl, re, sys
sys.stdout.reconfigure(encoding='utf-8')

p = r'C:\Users\Arhys\Desktop\каллибровка\фракционный состав\астм д 1160.xlsx'
wb = openpyxl.load_workbook(p, data_only=False)
ws = wb['Лист2']
print('--- MATCH-индексы в M/N по строкам ---')
for r in range(5, 16):
    m = ws['M%d' % r].value or ''
    n = ws['N%d' % r].value or ''
    l = ws['L%d' % r].value or ''
    mg = re.findall(r'MATCH\(Лист2!G(\d+)', m)
    ng = re.findall(r'MATCH\(Лист2!G(\d+)', n)
    if 'D$42' in l:
        lc = 'D'
    elif 'E$42' in l:
        lc = 'E'
    elif 'F$42' in l:
        lc = 'F'
    else:
        lc = '?'
    print('row %2d (B=%s): M MATCH G%s, N MATCH G%s, L->%s' % (r, ws['B%d' % r].value, mg, ng, lc))
wb.close()