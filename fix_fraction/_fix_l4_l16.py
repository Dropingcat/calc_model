# -*- coding: utf-8 -*-
"""Точечная доводка ФР-05: L4 (IBP) и L16 (FBP) в D1160 — переключение 1мм/10мм."""
import openpyxl, os

folder = r'C:\Users\Arhys\Desktop\каллибровка\фракционный состав'
p = os.path.join(folder, 'астм д 1160.xlsx')
wb = openpyxl.load_workbook(p, data_only=False)
ws = wb['Лист2']

# L4 (IBP): 1мм -> C40/C41/C42 (2.372/0/2.9); 10мм -> C44/C45/C46 (2.246/0/2.8)
ws['L4'] = "=IF('АСТМ Д 1160'!$E$1=Лист1!$N$2,Лист1!C42*(2.718281828*EXP(Лист1!C40+Лист1!C41*LN(1.8*Лист2!F4)))/1.8,Лист1!C46*(2.718281828*EXP((Лист1!C44+Лист1!C45*LN((1.8*Лист2!F4)))))/1.8)"

# L16 (FBP): 1мм -> F40/F41/F42 (0.718/0/2.9); 10мм -> F44/F45/F46 (1.521/0/2.8)
ws['L16'] = "=IF('АСТМ Д 1160'!$E$1=Лист1!$N$2,Лист1!F42*(2.718281828*EXP((Лист1!F40+Лист1!F41*LN((1.8*Лист2!G16)))))/1.8,Лист1!F46*(2.718281828*EXP((Лист1!F44+Лист1!F45*LN((1.8*Лист2!G16)))))/1.8)"

wb.save(p)

# verify
wb = openpyxl.load_workbook(p, data_only=False)
ws = wb['Лист2']
print('L4  =', str(ws['L4'].value)[:170])
print('L16 =', str(ws['L16'].value)[:170])
print('OK saved')